from docxtpl import DocxTemplate
import docx
from datetime import date
from use_db import get_book


def make_template(book_list, dir_name, name1, name2, name3, id):
    doc = DocxTemplate("act.docx")
    context = {'ФИО_руководителя': dir_name, 'ФИО1': name1, 'ФИО2': name2, 'ФИО3': name3, 'дата': date.today(),
               'количество': str(len(book_list))}
    doc.render(context)
    doc.save(f"act-final{str(id)}.docx")

    doc = docx.Document(f"act-final{str(id)}.docx")
    table = doc.add_table(rows=1, cols=len(book_list))

    for j in range(len(book_list)):
        cell = table.cell(j, 0)
        cell.text = get_book(book_list[j])

    doc.save(f"act-final{str(id)}.docx")
    return f"act-final{str(id)}.docx"

